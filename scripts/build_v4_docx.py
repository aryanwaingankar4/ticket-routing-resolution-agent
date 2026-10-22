"""Build Calibrated_Escalation_IT_Ticket_Triage_v4.docx from v3 (Phase 4).

Paper revision (Oct 2026). Offline, model-free, no LLM. Reads v3 (never
writes it), fills every [[FILL-...]] marker, applies the text changes the
authors approved (paper_docx/RESULTS_FOR_PAPER.md), inserts Figs. 4-6 at
3.3 in, adds a Tier-1 column to Table VIII, adds references [24]-[25] and a
new Appendix C, and writes v4.

Every number from the Phase 2 / Phase 3 evaluations is READ from its
committed result file and formatted here, never typed; the Phase 1 values
are asserted against their files. A missing marker, a target sentence that
is not found exactly once, or a value that does not match its source is
FATAL. Marker runs keep their font and size and lose the yellow shading and
the marker bold.

    .\\venv\\Scripts\\python.exe scripts/build_v4_docx.py [--force]
"""

from __future__ import annotations

import argparse
import copy
import csv
import json
import os
import re

from realworld_common import PROJECT_ROOT, die

DOCS = os.path.join(PROJECT_ROOT, "paper_docx")
V3 = os.path.join(DOCS, "Calibrated_Escalation_IT_Ticket_Triage_v3.docx")
V4 = os.path.join(DOCS, "Calibrated_Escalation_IT_Ticket_Triage_v4.docx")
FIG = os.path.join(DOCS, "fig")
DATA = os.path.join(PROJECT_ROOT, "data")
REV = os.path.join(DATA, "paper_revision_oct")
ZEN = os.path.join(DATA, "external_zenodo_7648117")
IND = os.path.join(DATA, "independent_test")
FIG_WIDTH_IN = 3.3
REF_PEREIRA, REF_LANGID = 24, 25


# ----------------------------------------------------------------- sources
def kv(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return {r["quantity"]: r["value"] for r in csv.DictReader(fh)}


def rows(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def check(cond, what):
    if not cond:
        die(f"source check failed: {what}")


def pct(k, n):
    return f"{100 * k / n:.1f}%"


def ci(s, key):
    return (f"[{100 * float(s[key + '.wilson95_lo']):.1f}, "
            f"{100 * float(s[key + '.wilson95_hi']):.1f}]")


def load_values():
    v = {}
    # Phase 2: real-world, primary mapping and sensitivity
    p = kv(os.path.join(ZEN, "realworld_summary.csv"))
    s = kv(os.path.join(ZEN, "realworld_summary_sens_fileservice_storage.csv"))
    check(p["mapping"] == "confirmed primary", "primary summary")
    check(s["mapping"] == "sensitivity:fileservice_storage", "sens summary")
    with open(os.path.join(ZEN, "category_mapping.json"), encoding="utf-8") as fh:
        check(json.load(fh)["status"] == "confirmed", "mapping confirmed")
    D = lambda k: int(p[f"distinct.{k}"])  # noqa: E731
    v.update(
        rw_n=D("n"), rw_ok=D("cascade_correct.k"), rw_t2=D("tier2_only_correct.k"),
        rw_esc=D("escalated.k"), rw_t1=D("tier1_kept.k"),
        rw_t1_ok=D("tier1_kept_correct.k"), rw_mis=D("misclassified"),
        rw_mis_draft=D("misclassified_drafted"),
        rw_mis_esc=D("misclassified_escalated"),
        rw_proc=D("proceeded.n"), rw_proc_ok=D("proceeded_correct.k"),
        rw_all_n=int(p["all.n"]), rw_all_ok=int(p["all.cascade_correct.k"]),
        rw_all_esc=int(p["all.escalated.k"]),
        rw_rec_am=p["distinct.recall.cascade.Access Management"],
        rw_rec_app=p["distinct.recall.cascade.Application"],
        rw_p=float(p["distinct.mcnemar_cascade_vs_tier2.p_exact"]),
        rw_excluded=int(p["excluded_total"]), rw_english=int(p["english_total"]),
        rw_ok_ci=ci(p, "distinct.cascade_correct"),
        rw_esc_ci=ci(p, "distinct.escalated"),
        rw_t1_ci=ci(p, "distinct.tier1_kept"),
        sens_ok=int(s["distinct.cascade_correct.k"]),
        sens_t2=int(s["distinct.tier2_only_correct.k"]),
        sens_ok_ci=ci(s, "distinct.cascade_correct"),
        sens_p=float(s["distinct.mcnemar_cascade_vs_tier2.p_exact"]),
        sens_rec_am=s["distinct.recall.cascade.Access Management"],
        sens_rec_app=s["distinct.recall.cascade.Application"],
        sens_rec_sto=s["distinct.recall.cascade.Storage"],
        sens_t1_ok=int(s["distinct.tier1_kept_correct.k"]),
        sens_proc_ok=int(s["distinct.proceeded_correct.k"]),
        sens_esc=int(s["distinct.escalated.k"]),
    )
    check(v["rw_mis_draft"] + v["rw_mis_esc"] == v["rw_mis"], "misrouted split")
    check(v["sens_esc"] == v["rw_esc"], "escalation is label-free")
    v["rw_ok_esc"] = v["rw_esc"] - v["rw_mis_esc"]
    v["rw_dup"] = v["rw_all_n"] - v["rw_n"]
    # Phase 2: draft labels -- the authors chose the "not labelled" wording
    with open(os.path.join(ZEN, "draft_labels.csv"), newline="",
              encoding="utf-8-sig") as fh:
        labelled = [r for r in csv.DictReader(fh)
                    if (r["grounded"] or r["appropriate_handoff"]).strip()]
    check(not labelled, "draft_labels.csv is unlabelled (paper says so)")
    with open(os.path.join(ZEN, "drafts.json"), encoding="utf-8") as fh:
        v["n_drafts"] = len(json.load(fh)["drafts"])
    # Phase 3: independent set
    i = kv(os.path.join(IND, "results", "independent_summary.csv"))
    check(i["annotation"] == "two authors, cross-labelled", "two annotators")
    v.update(
        ind_n=int(i["n_tickets"]), ind_ncat=int(i["n_with_category"]),
        ind_ok=int(i["with_category.cascade_correct.k"]),
        ind_t2=int(i["with_category.tier2_only_correct.k"]),
        ind_ok_ci=ci(i, "with_category.cascade_correct"),
        ind_esc=int(i["all.escalated.k"]), ind_esc_ci=ci(i, "all.escalated"),
        ind_out=int(i["all.outcome_correct.k"]),
        ind_out_ci=ci(i, "all.outcome_correct"),
        ind_se=int(i["should_escalate.escalated.k"]),
        ind_se_n=int(i["should_escalate.escalated.n"]),
        ind_sp=int(i["should_proceed.proceeded.k"]),
        ind_sp_n=int(i["should_proceed.proceeded.n"]),
        ind_t1=int(i["all.tier1_kept.k"]),
        ind_mis=int(i["with_category.misclassified"]),
        ind_mis_draft=int(i["with_category.misclassified_drafted"]),
        ind_rec_sec=i["with_category.recall.cascade.Security"],
        ind_rec_db=i["with_category.recall.cascade.Database"],
        k_cat=float(i["kappa.category"]),
        k_cat_ci=(f"[{float(i['kappa.category.boot95_lo']):.3f}, "
                  f"{float(i['kappa.category.boot95_hi']):.3f}]"),
        k_cat_agree=i["kappa.category.agreement"],
        k_out=float(i["kappa.outcome"]), k_out_agree=i["kappa.outcome.agreement"],
        ind_dis_cat=int(i["disagreements.category"]),
        ind_dis_out=int(i["disagreements.outcome"]),
        ind_adj=int(i["adjudicated"]),
        ind_adj_writer=int(i["adjudicated.final_equals_writer_label"]),
        ind_per_author=int(i["n_tickets.AW"]),
    )
    check(int(i["n_tickets.SP"]) == v["ind_per_author"], "54 per author")
    check(v["ind_adj"] == v["ind_dis_cat"] + v["ind_dis_out"], "adjudication")
    v["ind_noncat"] = v["ind_n"] - v["ind_ncat"]
    check(v["ind_se_n"] == v["ind_noncat"], "NONE rows are the escalate rows")
    # X1 says "seven per category plus five off-topic or vague" per author
    for a in ("AW", "SP"):
        comp = {}
        for r in rows(os.path.join(IND, f"tickets_{a}.csv")):
            comp[r["true_category"]] = comp.get(r["true_category"], 0) + 1
        check(comp.pop("NONE", 0) == 5 and len(comp) == 7
              and set(comp.values()) == {7}, f"composition of tickets_{a}")
    # X3 says recall ranges from Security (lowest) to Database (highest)
    rec = {k.rsplit(".", 1)[1]: int(val.split("/")[0])
           for k, val in i.items() if k.startswith("with_category.recall.cascade.")}
    check(min(rec, key=rec.get) == "Security" and rec["Security"] < min(
        x for c, x in rec.items() if c != "Security")
          and rec["Database"] > max(x for c, x in rec.items() if c != "Database"),
          "recall extremes are Security and Database")
    v["ind_act_esc"] = v["ind_sp_n"] - v["ind_sp"]
    # Phase 1 / benchmark values, asserted against their files
    e = kv(os.path.join(REV, "error_analysis_benchmark45_summary.csv"))
    check((e["n"], e["correct"], e["escalated"], e["misclassified"],
           e["misclassified_escalated"], e["misclassified_drafted"],
           e["tier1_kept"]) == ("45", "32", "15", "13", "5", "8", "4"),
          "benchmark-45 error-analysis summary")
    abl = [r for r in rows(os.path.join(DATA, "ablation_baseline_results.csv"))
           if r["section"] == "classification"]
    check(len(abl) == 45 and sum(r["tier_used"] == "1" for r in abl) == 4,
          "benchmark Tier-1 share 4/45")
    curve = {r["threshold"]: r for r in rows(os.path.join(
        REV, "cascade_calibration_curve_paraphrased175.csv"))}
    for th, kept, ok in (("0.50", "23", "19"), ("0.60", "10", "6"),
                         ("0.43", "35", "29"), ("0.36", "75", "53"),
                         ("0.46", "30", "25")):
        check((curve[th]["n_kept"], curve[th]["n_kept_correct"]) == (kept, ok),
              f"cascade curve at {th}")
    with open(os.path.join(DATA, "ood_calibration_tickets.json"),
              encoding="utf-8") as fh:
        ood = {t["id"]: t["text"] for t in json.load(fh)}
    check(ood["ood_D1_v3"].startswith("Hey there, something isn't working "
                                      "right on my end and I really need "
                                      "some assistance"), "E2 example text")
    return v


# ---------------------------------------------------------------- the text
def fills(v):
    rw_pct = pct(v["rw_ok"], v["rw_n"])
    rw_esc_pct = pct(v["rw_esc"], v["rw_n"])
    sens_pct = pct(v["sens_ok"], v["rw_n"])
    ind_pct = pct(v["ind_ok"], v["ind_ncat"])
    p_rw = "1.0" if v["rw_p"] == 1.0 else f"{v['rw_p']:.2f}"

    ind_x1 = (
        f"Two authors each wrote {v['ind_per_author']} tickets, seven per "
        f"category plus five off-topic or vague ones ({v['ind_n']} in total). "
        "They wrote without consulting the training data, the generation "
        "templates or the benchmarks, and without LLM assistance. Each author "
        "then labelled the category and the expected outcome (escalate or "
        "proceed) of the other's tickets without seeing the writer's labels. "
        f"Before adjudication, Cohen's κ was {v['k_cat']:.3f} for category "
        f"(bootstrap 95% CI {v['k_cat_ci']}; {v['k_cat_agree']} agreement) "
        f"and {v['k_out']:.3f} for outcome ({v['k_out_agree']}). The "
        f"{v['ind_dis_cat']} category disagreements were adjudicated, and in "
        f"all {v['ind_adj_writer']} the writer's label was kept.")
    check(v["ind_dis_out"] == 0 and v["ind_adj_writer"] == v["ind_adj"],
          "X1 wording assumes no outcome disagreement and writer labels kept")
    rw_x1 = (
        f"The real-world set is the public dataset of Pereira [{REF_PEREIRA}], "
        "released under CC BY 4.0: 2,229 tickets typed by the users of an "
        "IT-support company and labelled into seven categories by three IT "
        f"professionals. We kept the {v['rw_english']:,} tickets that "
        f"langid.py [{REF_LANGID}] identifies as English. Four source "
        "categories correspond to one of our teams, and the mapping was fixed "
        "before predictions were examined: Active Directory and Fileservice "
        "(almost entirely file-share access requests) to Access Management, "
        "and Software and O365 to Application. The other three categories "
        f"({v['rw_excluded']} tickets) have no single counterpart and were "
        "excluded, so Database, Infrastructure, Network, Security and Storage "
        f"are not covered. Of the {v['rw_all_n']} mapped tickets, "
        f"{v['rw_dup']} repeat an earlier text exactly, mostly "
        f"system-generated subject lines, so we report the {v['rw_n']} "
        f"distinct texts and give all {v['rw_all_n']} alongside. We found no "
        "public, real, English IT ticket dataset whose categories match our "
        "seven teams; the Zenodo corpus was the closest real option, and our "
        "hand-written set covers all seven.")

    ind_x3 = (
        f"On the independent set, the cascade routes {v['ind_ok']} of the "
        f"{v['ind_ncat']} tickets that belong to a team correctly ({ind_pct}, "
        f"95% CI {v['ind_ok_ci']}); Tier 2 alone routes "
        f"{v['ind_t2']}/{v['ind_ncat']}, and the out-of-template benchmark "
        "gave 32/45, so accuracy on this human-written text is similar to "
        "that on the Gemini-written benchmark. Per-team recall ranges from "
        f"{v['ind_rec_sec']} (Security) to {v['ind_rec_db']} (Database). The "
        f"retrieval gate escalates {v['ind_esc']} of the {v['ind_n']} tickets "
        f"({pct(v['ind_esc'], v['ind_n'])} {v['ind_esc_ci']}): "
        f"{v['ind_se']} of the {v['ind_se_n']} off-topic or vague tickets, but "
        f"also {v['ind_act_esc']} of the {v['ind_sp_n']} actionable ones, so "
        "the escalate-or-proceed decision is correct for "
        f"{v['ind_out']}/{v['ind_n']} ({pct(v['ind_out'], v['ind_n'])} "
        f"{v['ind_out_ci']}). Of the {v['ind_mis']} misrouted tickets, "
        f"{v['ind_mis_draft']} reached the drafter. Tier 1 answered "
        f"{v['ind_t1']} of the {v['ind_n']} tickets.")
    rw_x3 = (
        "On real tickets the classifier does not transfer: the cascade routes "
        f"{v['rw_ok']} of {v['rw_n']} distinct texts to the right team "
        f"({rw_pct}, 95% CI {v['rw_ok_ci']}; {v['rw_all_ok']}/{v['rw_all_n']}, "
        f"{pct(v['rw_all_ok'], v['rw_all_n'])}, over all tickets), with recall "
        f"{v['rw_rec_am']} for Access Management and {v['rw_rec_app']} for "
        f"Application. Tier 2 alone does no better ({v['rw_t2']}/{v['rw_n']}; "
        f"exact McNemar p = {p_rw}), so the cascade neither helps nor hurts "
        "accuracy here. Tier 1 answers far more often than on generated text "
        f"({v['rw_t1']}/{v['rw_n']}, {pct(v['rw_t1'], v['rw_n'])} "
        f"{v['rw_t1_ci']}, against 4/45 on the out-of-template set), but only "
        f"{v['rw_t1_ok']} of those {v['rw_t1']} answers are correct. The "
        f"retrieval gate escalates {v['rw_esc']}/{v['rw_n']} ({rw_esc_pct} "
        f"{v['rw_esc_ci']}; {v['rw_all_esc']}/{v['rw_all_n']}, "
        f"{pct(v['rw_all_esc'], v['rw_all_n'])}, over all tickets) but does "
        f"not filter misrouted tickets: {v['rw_mis_draft']} of the "
        f"{v['rw_mis']} misrouted texts still reached the drafter. The result "
        "also depends on one mapping decision: if the file-share access "
        "requests are mapped to Storage, where our generator files "
        f"share-permission faults, accuracy is {v['sens_ok']}/{v['rw_n']} "
        f"({sens_pct}, {v['sens_ok_ci']}; Appendix C), while the escalation "
        "rate and Tier-1 share do not depend on the mapping. This is the "
        "recurring finding of this paper in its plainest form: a classifier "
        "trained only on synthetic templates does not transfer to real "
        "tickets, and the gates, calibrated on the same synthetic material, "
        f"do not compensate for it. The {v['n_drafts']} drafts generated for a "
        "sample of the real tickets were not labelled.")

    # The whole abstract body (after "Abstract—"), shortened to <= 240 words
    # at the authors' request (2026-10-23); it replaces v3's body and A1.
    abstract = (
        "A triage system for IT support tickets must route each ticket to the "
        "right team and recognise when to hand it to a person. Our pipeline "
        "has two separately calibrated escalation gates: a 0.50 confidence "
        "threshold on a fast "
        "lexical classifier decides whether a slower embedding classifier is "
        "consulted, and a 0.67 cosine-similarity threshold on retrieved past "
        "tickets decides whether a language model may draft a resolution. "
        "On a 9-ticket adversarial set the "
        "retrieval gate decides every ticket correctly, against 3 of 9 "
        "without it. A zero-shot LLM classifies more accurately (40/45 "
        "against 33/45) but cannot abstain, so it serves as a baseline. "
        "On "
        f"{v['ind_n']} tickets written and cross-labelled by two authors "
        f"(κ = {v['k_cat']:.3f}), the cascade routes {v['ind_ok']} of "
        f"{v['ind_ncat']} correctly, but the retrieval gate also escalates "
        f"{v['ind_act_esc']} actionable tickets. On {v['rw_n']} distinct real "
        "tickets from a public dataset, covering two of seven teams, it "
        f"routes {rw_pct} correctly "
        f"({pct(v['rw_all_ok'], v['rw_all_n'])} over all {v['rw_all_n']} "
        "tickets), indistinguishable from Tier 2 alone "
        f"(p = {p_rw}), or {sens_pct} under an alternative category mapping; "
        f"Tier 1 answers {v['rw_t1']} of them with {v['rw_t1_ok']} correct, "
        f"and {v['rw_mis_draft']} of {v['rw_mis']} misrouted tickets still "
        "reach the drafter. One pattern repeats across experiments: the data "
        "a component is "
        "trained or calibrated on limits it more than the method chosen, and "
        "a classifier trained only on synthetic templates does not transfer "
        "to real tickets. We report negative results, an error analysis and "
        "reproduction steps; conclusions are scoped to small evaluation "
        "sets.")
    if len(abstract.split()) > 240:
        die(f"abstract is {len(abstract.split())} words (limit 240)")
    # Sentences the new sets made stale (authors' request, 2026-10-23)
    stale = [
        ("Two of the ticket-triage studies evaluate on real tickets, while our "
         "main evaluation does not (Section VIII).",
         "Two of the ticket-triage studies evaluate on real tickets, while our "
         f"main evaluation is synthetic, supplemented by {v['ind_n']} "
         f"hand-written tickets and {v['rw_n']} real tickets that cover only "
         "two of our seven teams (Section VI-E)."),
        ("No real ticket stream has passed through the system and",
         f"The system has not been deployed: {v['rw_n']} real tickets were "
         "routed offline, but no live ticket stream has passed through it "
         "and"),
        ("These conclusions are established on a synthetic corpus and small, "
         "mostly LLM-generated evaluation sets, and they",
         "These conclusions are established on a synthetic corpus, small "
         f"LLM-generated evaluation sets, {v['ind_n']} hand-written tickets "
         f"and {v['rw_n']} real tickets that cover only two of the seven "
         "teams, and they"),
    ]
    c1 = (
        f"On the independent author-written set, routing accuracy "
        f"({v['ind_ok']}/{v['ind_ncat']}) was similar to that on the "
        "out-of-template benchmark, but the retrieval gate escalated "
        f"{v['ind_act_esc']} of the {v['ind_sp_n']} actionable tickets. On "
        "real user-written tickets for the two teams we could test, the "
        f"classifier did not transfer: {v['rw_ok']} of {v['rw_n']} distinct "
        f"texts ({rw_pct}; {v['rw_all_ok']}/{v['rw_all_n']} over all tickets) "
        f"were routed correctly, Tier 2 alone did no better (p = {p_rw}), and "
        f"Tier 1 answered {v['rw_t1']} of them with only {v['rw_t1_ok']} "
        "correct. The retrieval gate did not filter the misroutes "
        f"({v['rw_mis_draft']} of {v['rw_mis']} reached the drafter), and "
        f"accuracy depends on the category mapping ({sens_pct} under the "
        "alternative). This is the paper's recurring finding at its starkest: "
        "the synthetic training data, not the choice of method or threshold, "
        "limits what the pipeline can do.")
    l1 = (
        "The independent set removes the dependence on LLM-written evaluation "
        f"text for all seven teams, but its {v['ind_n']} tickets were written "
        "and labelled by two of the authors, who know the categories, so it "
        "is still small and is not real traffic. The real-world set removes "
        "that dependence for two teams only, because no public, real, English "
        "IT ticket dataset with categories matching our seven teams was "
        "available: Database, Infrastructure, Network, Security and Storage "
        "remain untested on real tickets. Its accuracy depends strongly on "
        f"how the source categories are mapped to ours ({rw_pct} or "
        f"{sens_pct}), and the tickets come from a single organisation. The "
        "English filter also let a few non-English tickets through, because "
        "their system-generated subject lines were in English.")
    e4 = (
        "of the 13 misclassified tickets, 5 were escalated by the retrieval "
        "gate (the safe outcome) and 8 received a draft (the unsafe outcome). "
        "The classification gate does not stop a misrouted ticket from "
        "reaching the drafter, which is a limitation of the current design. "
        f"The new sets show the same at larger scale: {v['ind_mis_draft']} of "
        f"the {v['ind_mis']} misrouted independent tickets reached the "
        f"drafter, and of the {v['rw_mis']} misrouted distinct real texts, "
        f"{v['rw_mis_draft']} passed the retrieval gate and reached the "
        f"drafter. The gate escalated {v['rw_mis_esc']} misrouted and "
        f"{v['rw_ok_esc']} correctly routed real texts, so its similarity "
        "score does not track routing correctness. Tier 1's confidence does "
        f"not either: it kept {v['rw_t1']} of the {v['rw_n']} real texts at "
        f"the 0.50 threshold, and {v['rw_t1_ok']} of those were correct")
    t1 = (
        "The target accuracies were 0.90, 0.80 and 0.70. No threshold met "
        "0.90, so that target would have sent every ticket to Tier 2. At "
        "0.50, Tier 1 keeps 23 of the 175 paraphrased tickets (13.1%) and "
        "labels 19 of them correctly (82.6%). The curve is not monotone: "
        "above 0.60 only 10 tickets remain, of which 6 are correct, so "
        "accuracy falls again at small counts. The rule takes the lowest "
        "bucket edge whose cumulative accuracy meets the target, which is "
        "0.50. On a 0.01 grid the same rule would choose 0.43 for the 0.80 "
        "target and 0.36 for 0.70; kept accuracy peaks at 25/30 at 0.46, so "
        "0.50 sits near the peak.")
    return {
        "ABSTRACT": abstract, "STALE": stale,
        "C1": c1, "L1": l1, "E4": e4, "T1": t1,
        "X1": ind_x1 + " " + rw_x1, "X3": ind_x3 + " " + rw_x3,
        "D1": ("The seven categories are Access Management, Application, "
               "Database, Infrastructure, Network, Security and Storage."),
        "D2": (f"We add two evaluation sets that no language model wrote: "
               f"{v['ind_n']} tickets written by two of the authors, each "
               "labelling the other's tickets blind, and "
               f"{v['rw_n']} distinct real tickets from a public IT-support "
               "dataset (Section VI-E)."),
        "T0": (f"independent author-written set ({v['ind_n']} tickets); "
               f"{v['rw_n']} real tickets [{REF_PEREIRA}]"),
        "S1": "800", "S2": "175", "S4": "45",
        "S5": "Gemini variants of 15 hand-written off-topic and vague seeds",
        "S6": str(v["ind_n"]),
        "S7": f"{v['rw_n']} distinct ({v['rw_all_n']})",
        "S8": f"Real user tickets, Pereira [{REF_PEREIRA}], CC BY 4.0",
        "H1": "unigrams and bigrams, English stop words removed",
        "H2": ("max_iter = 1000; otherwise scikit-learn defaults (L2, C = 1.0, "
               "lbfgs)"),
        "H3": ("max_iter = 1000, scikit-learn defaults, fitted on the "
               "3,200-row split"),
        "H4": ("API defaults (no generation config set); zero-shot runs at "
               "temperature 0.0"),
        "H5": ("4 epochs, learning rate 2 × 10⁻⁵, batch size 16, maximum "
               "length 128, weight decay 0.01"),
        "H7": ("seed 42 throughout (stratified 80/20 split); latency on an "
               "Intel Core i5-1334U (12 logical CPUs), Windows 11, Python "
               "3.14.3, median of 200 runs after 20 warm-up runs"),
        "X2": "15/45 (33.3%)",
        "X4": str(v["ind_n"]), "X5": f"{v['ind_ok']}/{v['ind_ncat']}",
        "X6": f"{v['ind_t2']}/{v['ind_ncat']}",
        "X7": f"{v['ind_esc']}/{v['ind_n']} ({pct(v['ind_esc'], v['ind_n'])})",
        "X8": str(v["rw_n"]), "X9": f"{v['rw_ok']}/{v['rw_n']}",
        "X10": f"{v['rw_t2']}/{v['rw_n']}",
        "X11": f"{v['rw_esc']}/{v['rw_n']} ({rw_esc_pct})",
        "E1": ("Database→Application, Infrastructure→Application and "
               "Security→Storage (2 each); 10 of the 13 errors are wrong "
               "predictions of Application or Storage"),
        "E2": ("all 6 are vague ‘something is broken’ messages naming no "
               "system (3 of the 15 seeds); no off-topic ticket (weather, "
               "food, HR, Finance, Facilities) leaks. Example: ‘Hey there, "
               "something isn't working right on my end and I really need "
               "some assistance…’ (0.725)"),
        "E3": ("For example, N31 (Security: ransomware pop-ups demanding "
               "payment) was routed to Storage, N39 (Network: the VPN "
               "connection to the company portal fails) to Database, and N13 "
               "(Infrastructure: company websites down for the whole office) "
               "to Database."),
        "E5": ("In both drafts the model itself noted that the retrieved "
               "tickets did not match. N26 retrieved fixes for failing disks "
               "behind network file shares, and N31 retrieved NTFS permission "
               "fixes. It then supplied steps from outside knowledge, which "
               "the prompt forbids: clean up the local disk for N26, and "
               "disconnect the machine and escalate to incident response for "
               "N31. Both drafts asked for human review. N31 had also been "
               "misrouted to Storage."),
    }, v


# Approved non-marker edits: (unique old text, new text)
TEXT_EDITS = [
    # 1b, §VI-A
    ("(two standard deviations over calibration draws)",
     "(two standard deviations of a binomial coverage estimate, n = 175)"),
    # 3a, §IV-C
    ("The classifiers are trained on 3,200 rows.",
     "Tier 2 and DistilBERT are trained on a stratified 3,200-row split "
     "(seed 42); the production Tier 1 is fitted on all 4,000 rows."),
    # 4, §V-A
    ("chose the lowest threshold that meets the target accuracy.",
     "chose the lowest threshold that meets the target accuracy on 10%-wide "
     "confidence buckets."),
    # 1c, Table VI caption
    ("Only the calibration set changes between the two gap columns",
     "Only the calibration set changes between the two gap columns. The "
     "±2 s.d. column is twice the binomial standard deviation of a coverage "
     "estimate at n = 175"),
]


# ------------------------------------------------------------ docx helpers
def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()
    if os.path.abspath(V4) == os.path.abspath(V3):
        die("refusing: output path is v3")
    if os.path.exists(V4) and not args.force:
        die(f"{V4} exists (use --force to rebuild it; v3 is never touched)")

    import docx
    from docx.oxml.ns import qn
    from docx.shared import Inches
    from docx.text.paragraph import Paragraph

    F, v = fills(load_values())
    d = docx.Document(V3)
    body = d.element.body
    W = lambda tag: qn("w:" + tag)  # noqa: E731

    def text_runs(p):
        out = []
        for r in p.iter(W("r")):
            ts = r.findall(W("t"))
            if ts:
                out.append((r, ts))
        return out

    def is_marker_run(r):
        rpr = r.find(W("rPr"))
        shd = rpr.find(W("shd")) if rpr is not None else None
        return shd is not None and shd.get(qn("w:fill")) == "FFFF00"

    def unmark(r):
        rpr = r.find(W("rPr"))
        for tag in ("b", "bCs", "shd", "highlight"):
            for x in rpr.findall(W(tag)):
                rpr.remove(x)

    def set_text(ts, text):
        ts[0].text = text
        ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        for t in ts[1:]:
            t.getparent().remove(t)

    def replace_in(p, old, new):
        """Replace `old` (exactly once in p) with `new`, across runs."""
        tr = text_runs(p)
        texts = ["".join(t.text or "" for t in ts) for _, ts in tr]
        full = "".join(texts)
        if full.count(old) != 1:
            die(f"expected exactly one {old[:60]!r}, found {full.count(old)}")
        start, end = full.index(old), full.index(old) + len(old)
        pos, a, b = 0, None, None
        offs = []
        for k, t in enumerate(texts):
            offs.append(pos)
            if a is None and start < pos + len(t):
                a = k
            if end <= pos + len(t) and b is None and a is not None:
                b = k
            pos += len(t)
        pre = texts[a][:start - offs[a]]
        post = texts[b][end - offs[b]:]
        marker = is_marker_run(tr[a][0])
        if a == b:
            set_text(tr[a][1], pre + new + post)
        else:
            set_text(tr[a][1], pre + new)
            for k in range(a + 1, b):
                tr[k][0].getparent().remove(tr[k][0])
            if post:
                set_text(tr[b][1], post)
                if is_marker_run(tr[b][0]):
                    unmark(tr[b][0])
            else:
                tr[b][0].getparent().remove(tr[b][0])
        if marker:
            unmark(tr[a][0])

    def all_paras():
        return list(body.iter(W("p")))

    def ptext(p):
        return "".join(t.text or "" for t in p.iter(W("t")))

    def find_para(needle):
        hits = [p for p in all_paras() if needle in ptext(p)]
        if len(hits) != 1:
            die(f"expected one paragraph containing {needle[:60]!r}, "
                f"found {len(hits)}")
        return hits[0]

    marker_re = re.compile(r"\[\[FILL-([A-Z]+\d+)(?::[^\]]*)?\]\]")

    def marker_para(mid):
        hits = [p for p in all_paras()
                if mid in [x.group(1) for x in marker_re.finditer(ptext(p))]]
        if len(hits) != 1:
            die(f"marker {mid}: expected one paragraph, found {len(hits)}")
        return hits[0]

    def fill_marker(mid, text):
        p = marker_para(mid)
        m = [x for x in marker_re.finditer(ptext(p)) if x.group(1) == mid]
        if len(m) != 1:
            die(f"marker {mid} not found exactly once")
        replace_in(p, m[0].group(0), text)
        return p

    def clone_para(template, text, size=None, jc=None, first_line=None,
                   bold_prefix=None):
        """New paragraph with template's pPr and its first run's rPr."""
        p = copy.deepcopy(template)
        for r in list(p.iter(W("r"))):
            r.getparent().remove(r)
        for x in list(p):
            if x.tag not in (W("pPr"),):
                p.remove(x)
        ppr = p.find(W("pPr"))
        if ppr is not None:
            for x in ppr.findall(W("pBdr")):
                ppr.remove(x)
            # Schema order: ... spacing, ind, ..., jc, ..., rPr. Set in place
            # where the element exists; otherwise insert before what follows.
            if first_line is not None:
                ind = ppr.find(W("ind"))
                if ind is None:
                    ind = ppr.makeelement(W("ind"), {})
                    nxt = next((ppr.find(W(t)) for t in ("jc", "rPr")
                                if ppr.find(W(t)) is not None), None)
                    (nxt.addprevious(ind) if nxt is not None
                     else ppr.append(ind))
                for k in list(ind.attrib):
                    del ind.attrib[k]
                ind.set(qn("w:firstLine"), str(first_line))
            if jc is not None:
                el = ppr.find(W("jc"))
                if el is None:
                    el = ppr.makeelement(W("jc"), {})
                    rp = ppr.find(W("rPr"))
                    (rp.addprevious(el) if rp is not None else ppr.append(el))
                el.set(qn("w:val"), jc)
        src = template.find(".//" + W("r"))
        rpr = copy.deepcopy(src.find(W("rPr")))
        for tag in ("b", "bCs", "shd", "highlight", "smallCaps"):
            for x in rpr.findall(W(tag)):
                rpr.remove(x)
        if size is not None:
            for tag in ("sz", "szCs"):
                for x in rpr.findall(W(tag)):
                    x.set(qn("w:val"), str(size))

        def add_run(t, bold=False):
            r = p.makeelement(W("r"), {})
            rp = copy.deepcopy(rpr)
            if bold:
                rp.insert(1, rp.makeelement(W("b"), {}))
            r.append(rp)
            te = r.makeelement(W("t"), {})
            te.text = t
            te.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            r.append(te)
            p.append(r)
        if bold_prefix:
            add_run(bold_prefix, bold=True)
        add_run(text)
        return p

    # 1. markers -------------------------------------------------------------
    # Abstract: replace the whole body. Its first run is bold body text, so
    # the new text keeps the abstract's bold (an A1 fill on its own would
    # inherit the marker run, whose bold is stripped).
    p = marker_para("A1")
    body_old = ptext(p)[ptext(p).index("Automatic triage"):]
    if not body_old.endswith("all conclusions are scoped accordingly."):
        die("v3 abstract body not as expected")
    replace_in(p, body_old, F["ABSTRACT"])
    for mid in ("D1", "D2", "T0", "S1", "S4", "S5", "S6", "S7", "S8",
                "H1", "H2", "H3", "H4", "H5", "H7", "T1", "X1", "X2", "X3",
                "X4", "X5", "X6", "X7", "X8", "X9", "X10", "X11", "E1", "E2",
                "E3", "E4", "E5", "L1", "C1"):
        fill_marker(mid, F[mid])

    # H6: the approved cell rewording replaces the marker and its frame
    p = marker_para("H6")
    m = marker_re.search(ptext(p))
    replace_in(p, f" over {m.group(0)} calibration draws",
               " of a binomial coverage estimate, n = 175")

    # Table III: merged calibration row (S2), deleted S3 row, training row,
    # independent-set source
    p = marker_para("S2")
    row = p.getparent().getparent()          # p -> tc -> tr
    cells = row.findall(W("tc"))
    fill_marker("S2", F["S2"])
    replace_in(cells[2].find(W("p")), "Paraphrased tickets",
               "Gemini paraphrases of training rows")
    replace_in(cells[3].find(W("p")), "Cascade threshold (0.50)",
               "Cascade threshold (0.50), conformal calibration, drift "
               "reference, in-domain arm of the 0.67 sweep")
    p = marker_para("S3")
    tr = p.getparent().getparent()
    if "In-domain calibration" not in ptext(tr):
        die("S3 row is not the In-domain calibration row")
    tr.getparent().remove(tr)
    p = find_para("Fitting all trained classifiers")
    tr = p.getparent().getparent()
    tcs = tr.findall(W("tc"))
    if ptext(tcs[0]) != "Training" or ptext(tcs[1]) != "3,200":
        die("Table III training row not as expected")
    replace_in(tcs[1].find(W("p")), "3,200", "3,200 / 4,000")
    replace_in(p, "Fitting all trained classifiers",
               "Tier 2 and DistilBERT on the 3,200-row split; Tier 1 on all "
               "4,000 rows")
    replace_in(find_para("Written blind by the authors"),
               "Written blind by the authors",
               "Written blind by two authors; cross-labelled")

    # 2. approved text edits -------------------------------------------------
    for old, new in TEXT_EDITS + F["STALE"]:
        replace_in(find_para(old), old, new)

    # 3. Table VIII: Tier-1 share column ---------------------------------------
    tbl = None
    for t in body.iter(W("tbl")):
        if "Out-of-template (ref.)" in ptext(t):
            tbl = t
    if tbl is None:
        die("Table VIII not found")
    widths = [1150, 500, 800, 800, 900, 700]
    grid = tbl.find(W("tblGrid"))
    gcs = grid.findall(W("gridCol"))
    if len(gcs) != 5 or sum(int(g.get(qn("w:w"))) for g in gcs) != sum(widths):
        die("Table VIII grid not as expected")
    new_gc = copy.deepcopy(gcs[-1])
    grid.append(new_gc)
    for g, w in zip(grid.findall(W("gridCol")), widths):
        g.set(qn("w:w"), str(w))
    values = {"Set": "Tier-1 share", "Out-of-template (ref.)": "4/45",
              "Independent": f"{v['ind_t1']}/{v['ind_n']}",
              "Real-world": f"{v['rw_t1']}/{v['rw_n']}"}
    for tr in tbl.findall(W("tr")):
        tcs = tr.findall(W("tc"))
        key = ptext(tcs[0])
        new = copy.deepcopy(tcs[-1])
        tcs[-1].addnext(new)
        tr_tp = new.find(".//" + W("p"))
        tr_runs = text_runs(tr_tp)
        set_text(tr_runs[0][1], values[key])
        for r, _ in tr_runs[1:]:
            r.getparent().remove(r)
        for tc, w in zip(tr.findall(W("tc")), widths):
            tc.find(W("tcPr")).find(W("tcW")).set(qn("w:w"), str(w))

    # 4. figures ------------------------------------------------------------
    fig1 = next(p for p in all_paras() if p.findall(".//" + W("drawing")))
    for mid, fname in (("F4", "fig4_cascade_threshold.png"),
                       ("F5", "fig5_retrieval_sweep.png"),
                       ("F6", "fig6_confusion_matrix.png")):
        p = marker_para(mid)
        for r in list(p.iter(W("r"))):
            r.getparent().remove(r)
        p.remove(p.find(W("pPr")))
        p.insert(0, copy.deepcopy(fig1.find(W("pPr"))))
        Paragraph(p, d._body).add_run().add_picture(
            os.path.join(FIG, fname), width=Inches(FIG_WIDTH_IN))

    # 5. prompts (Appendix A) -------------------------------------------------
    with open(os.path.join(REV, "prompts.txt"), encoding="utf-8") as fh:
        raw = fh.read()
    p1_raw, p2_raw = raw.split("=== P2:")
    p1_lines = p1_raw.splitlines()[3:]        # header, note, blank
    p2_lines = p2_raw.splitlines()[2:]        # rest of header line, blank
    while p1_lines and not p1_lines[-1].strip():
        p1_lines.pop()
    for mid, note, lines in (
            ("P1", "The ‘Past ticket #k’ block is repeated for each of the "
                   "top-5 retrieved tickets, k = 1…5.", p1_lines),
            ("P2", "", p2_lines)):
        p = fill_marker(mid, note)
        anchor = p
        for line in lines:
            q = clone_para(p, line, size=16, jc="left", first_line=0)
            anchor.addnext(q)
            anchor = q

    # 6. reproduction steps (Appendix B) --------------------------------------
    steps = [
        ("Environment: ", "python -m venv venv; .\\venv\\Scripts\\Activate.ps1; "
         "pip install -r requirements.txt (Python 3.14.3); for the revision "
         "analyses also pip install -r requirements-revision.txt. "
         "GEMINI_API_KEY in .env is needed only for steps that call Gemini."),
        ("Data, models and index: ", "python data/generate_dataset.py; "
         "python src/classification/train_tier1.py; python "
         "src/classification/train_embeddings.py; python "
         "src/rag/build_vector_index.py."),
        ("Threshold calibration: ", "python src/classification/train_cascade.py; "
         "python -m src.experiments.calibrate_rag_similarity_threshold; "
         "python -m src.experiments.calibrate_conformal."),
        ("Evaluation: ", "python src/experiments/test_adversarial_escalation.py; "
         "python src/experiments/run_ablation_study.py --mode {baseline, "
         "tier2-only, no-cascade, no-rag}; python "
         "src/experiments/compare_cascade_vs_tier2.py; python "
         "src/experiments/measure_inference_latency.py; python "
         "src/experiments/run_zeroshot_baselines.py and "
         "compare_zeroshot_vs_tier2.py; python "
         "src/experiments/score_groundedness_set.py; python "
         "src/experiments/score_sufficiency_gate.py."),
        ("Revision analyses and figures: ", "python "
         "scripts/reproduce_cascade_calibration_curve.py; python "
         "scripts/reproduce_ood_similarities.py; python "
         "scripts/error_analysis_benchmark45.py; python "
         "scripts/export_prompts.py; python scripts/make_revision_figures.py; "
         "python scripts/make_paper_figures.py."),
        ("Extended evaluation: ", "python scripts/fetch_zenodo_tickets.py; "
         "python scripts/prepare_zenodo_tickets.py; python "
         "scripts/run_realworld_routing.py; python scripts/score_realworld.py "
         "(and --variant fileservice_storage); python "
         "scripts/draft_realworld.py (calls Gemini); python "
         "scripts/eval_independent.py."),
        ("Tables, numbers and checks: ", "python "
         "src/experiments/build_paper_artifacts.py --force; pytest."),
    ]
    for _, cmd in steps:
        for path in re.findall(r"(?:src|scripts|data)/[\w/]+\.py", cmd):
            if not os.path.exists(os.path.join(PROJECT_ROOT, path)):
                die(f"R1 names a missing script: {path}")
        for mod in (x.rstrip(".") for x in re.findall(r"-m (src\.[\w.]+)", cmd)):
            if not os.path.exists(os.path.join(PROJECT_ROOT,
                                               *mod.split(".")) + ".py"):
                die(f"R1 names a missing module: {mod}")
    p = marker_para("R1")
    template = copy.deepcopy(p)
    fill_marker("R1", "")
    for r in list(p.iter(W("r"))):
        r.getparent().remove(r)
    first = clone_para(template, steps[0][1], jc="left", first_line=0,
                       bold_prefix="1. " + steps[0][0])
    p.addprevious(first)
    p.getparent().remove(p)
    anchor = first
    for k, (head, cmd) in enumerate(steps[1:], start=2):
        q = clone_para(template, cmd, jc="left", first_line=0,
                       bold_prefix=f"{k}. {head}")
        anchor.addnext(q)
        anchor = q
    last_b = anchor

    # 7. references [24], [25] ------------------------------------------------
    refs = [p for p in all_paras()
            if p.find(W("pPr")) is not None
            and p.find(W("pPr")).find(W("numPr")) is not None
            and p.find(W("pPr")).find(W("numPr")).find(W("numId")).get(
                qn("w:val")) == "4"]
    if len(refs) != REF_PEREIRA - 1:
        die(f"expected {REF_PEREIRA - 1} references, found {len(refs)}")
    new_refs = [
        "L. S. B. Pereira, “Classification of IT support tickets,” Zenodo, "
        "Dec. 2022, doi: 10.5281/zenodo.7648117. Dataset, CC BY 4.0.",
        "M. Lui and T. Baldwin, “langid.py: An off-the-shelf language "
        "identification tool,” in Proc. ACL 2012 System Demonstrations, "
        "Jeju Island, Korea, 2012, pp. 25–30.",
    ]
    anchor = refs[-1]
    for text in new_refs:
        q = clone_para(refs[-1], text)
        anchor.addnext(q)
        anchor = q

    # 8. Appendix C -------------------------------------------------------------
    head_b = find_para("Appendix B: Reproduction Steps")
    cap_no = find_para("TABLE IX")
    cap_txt = find_para("Error taxonomy of the live pipeline")
    t9 = cap_txt.getnext()
    if t9.tag != W("tbl"):
        die("Table IX does not follow its caption")
    head_c = copy.deepcopy(head_b)
    replace_in(head_c, "Appendix B: Reproduction Steps",
               "Appendix C: Sensitivity of the Real-World Result to the "
               "Category Mapping")
    intro = clone_para(last_b, (
        "Table X repeats the real-world evaluation of Section VI-E with the "
        "Fileservice category mapped to Storage instead of Access Management; "
        "everything else is unchanged. Sources: realworld_summary.csv and "
        "realworld_summary_sens_fileservice_storage.csv in "
        "data/external_zenodo_7648117/."), jc="left")
    cap1 = copy.deepcopy(cap_no)
    replace_in(cap1, "TABLE IX", "TABLE X")
    cap2 = copy.deepcopy(cap_txt)
    replace_in(cap2, "Error taxonomy of the live pipeline",
               "Real-world results under the primary mapping (file-share "
               "access requests → Access Management) and the sensitivity "
               f"mapping (→ Storage). Distinct texts, n = {v['rw_n']}. 95% "
               "Wilson intervals in brackets")
    rw_pct, sens_pct = pct(v["rw_ok"], v["rw_n"]), pct(v["sens_ok"], v["rw_n"])
    p_s = f"{v['sens_p']:.2f}"
    body_rows = [
        ("Quantity", "Primary", "Sensitivity"),
        ("Cascade correct", f"{v['rw_ok']}/{v['rw_n']} ({rw_pct}) {v['rw_ok_ci']}",
         f"{v['sens_ok']}/{v['rw_n']} ({sens_pct}) {v['sens_ok_ci']}"),
        ("Tier 2 alone correct", f"{v['rw_t2']}/{v['rw_n']}",
         f"{v['sens_t2']}/{v['rw_n']}"),
        ("Exact McNemar, cascade vs Tier 2",
         f"p = {'1.0' if v['rw_p'] == 1.0 else v['rw_p']}", f"p = {p_s}"),
        ("Recall, Access Management", v["rw_rec_am"], v["sens_rec_am"]),
        ("Recall, Application", v["rw_rec_app"], v["sens_rec_app"]),
        ("Recall, Storage", "–", v["sens_rec_sto"]),
        ("Tier-1 answers correct", f"{v['rw_t1_ok']}/{v['rw_t1']}",
         f"{v['sens_t1_ok']}/{v['rw_t1']}"),
        ("Drafted and correctly routed", f"{v['rw_proc_ok']}/{v['rw_proc']}",
         f"{v['sens_proc_ok']}/{v['rw_proc']}"),
        ("Escalated at 0.67",
         f"{v['rw_esc']}/{v['rw_n']} ({pct(v['rw_esc'], v['rw_n'])})",
         f"{v['sens_esc']}/{v['rw_n']} ({pct(v['sens_esc'], v['rw_n'])})"),
    ]
    tc_tbl = copy.deepcopy(t9)
    trs = tc_tbl.findall(W("tr"))
    hdr, mid_row, last_row = trs[0], trs[1], trs[-1]
    for tr in trs:
        tc_tbl.remove(tr)
    widths = [1950, 1450, 1450]
    for g, w in zip(tc_tbl.find(W("tblGrid")).findall(W("gridCol")), widths):
        g.set(qn("w:w"), str(w))
    for k, cells in enumerate(body_rows):
        src = hdr if k == 0 else (last_row if k == len(body_rows) - 1
                                  else mid_row)
        tr = copy.deepcopy(src)
        for tc, text, w in zip(tr.findall(W("tc")), cells, widths):
            tc.find(W("tcPr")).find(W("tcW")).set(qn("w:w"), str(w))
            tp = tc.find(W("p"))
            trn = text_runs(tp)
            set_text(trn[0][1], text)
            if is_marker_run(trn[0][0]):
                unmark(trn[0][0])
            for r, _ in trn[1:]:
                r.getparent().remove(r)
        tc_tbl.append(tr)
    anchor = last_b
    for el in (head_c, intro, cap1, cap2, tc_tbl):
        anchor.addnext(el)
        anchor = el

    # 8b. Length cuts approved by AW (2026-10-23): cuts 1, 4, 7, 8 (fold),
    # 9. Moved material keeps its content and its table number; the text
    # that cites it now names the appendix. Cuts 2, 3, 5 and 6 are NOT
    # applied.
    def appendix_heading(title):
        h = copy.deepcopy(head_b)
        replace_in(h, "Appendix B: Reproduction Steps", title)
        return h

    def table_block(caption_no):
        """Caption number, caption text, table and its spacer paragraph."""
        caps = [p for p in all_paras() if ptext(p).strip() == caption_no]
        if len(caps) != 1:
            die(f"{caption_no}: expected one caption, found {len(caps)}")
        cap = caps[0]
        txt = cap.getnext()
        tbl_el = txt.getnext()
        spacer = tbl_el.getnext()
        if tbl_el.tag != W("tbl") or spacer.tag != W("p") or ptext(spacer):
            die(f"{caption_no} block not as expected")
        return [cap, txt, tbl_el, spacer]

    # Cut 1: Table IV -> Appendix D
    block_iv = table_block("TABLE IV")
    old = "Table IV gives the configuration of every component."
    replace_in(find_para(old), old, "Table IV (Appendix D) gives the "
               "configuration of every component.")
    # Cut 4: Table II -> Appendix E; §III-G kept as two sentences
    block_ii = table_block("TABLE II")
    old = ("A value that is wrong for its context but internally consistent "
           "can produce a wrong result without raising any error. The checks "
           "in Table II therefore run before any result is accepted. Routing "
           "decisions are compared exactly, and floating-point values are "
           "compared within tolerances derived from their arithmetic.")
    replace_in(find_para(old), old,
               "A value that is wrong for its context but internally "
               "consistent can produce a wrong result without raising any "
               "error, so the checks in Table II (Appendix E) run before any "
               "result is accepted. Routing decisions are compared exactly, "
               "and floating-point values within tolerances derived from "
               "their arithmetic.")
    # Cut 7: "Methodological lessons" paragraph -> Appendix F, with a pointer
    lessons = find_para("Methodological lessons.")
    if not ptext(lessons).startswith("Methodological lessons."):
        die("Methodological lessons paragraph not as expected")
    old = ("training on the corpus does not transfer to out-of-template "
           "phrasing better than no training at all (Section VI-C).")
    replace_in(find_para(old), old, old + " Methodological lessons from "
               "these experiments are summarised in Appendix F.")
    # Cut 8: fold §III-F into §III-A as one sentence; reletter G -> F
    head_f = find_para("F. Automation candidates")
    para_f = head_f.getnext()
    fold = ("An offline job clusters resolved tickets by the similarity of "
            "their resolution text at a threshold of 0.80 and shows recurring "
            "clusters to a reviewer as candidates for automation.")
    if ptext(para_f) != fold or ptext(head_f) != "F. Automation candidates":
        die("§III-F not as expected")
    old = "because a failed ticket needs a person rather than a retry."
    replace_in(find_para(old), old, old + " An offline job also clusters "
               "resolved tickets by the similarity of their resolution text "
               "at a threshold of 0.80 and shows recurring clusters to a "
               "reviewer as candidates for automation.")
    for el in (head_f, para_f):
        el.getparent().remove(el)
    head_g = find_para("G. Verification")
    if ptext(head_g) != "G. Verification":
        die("§III-G heading not as expected")
    replace_in(head_g, "G. Verification", "F. Verification")
    # Cut 9: Fig. 1 at 85% (both dimensions, aspect kept)
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    for tag in (WP + "extent", A + "ext"):
        for ext in fig1.iter(tag):
            for k in ("cx", "cy"):
                ext.set(k, str(round(int(ext.get(k)) * 0.85)))
    # Keep Table III in one column: with the cuts it would otherwise break
    # after its first row. keepNext on every row but the last.
    t3 = [t for t in body.iter(W("tbl"))
          if "Paraphrased calibration" in ptext(t)]
    if len(t3) != 1:
        die("Table III not found")
    for tr in t3[0].findall(W("tr"))[:-1]:
        for tp in tr.iter(W("p")):
            ppr = tp.find(W("pPr"))
            if ppr is None:
                ppr = tp.makeelement(W("pPr"), {})
                tp.insert(0, ppr)
            if ppr.find(W("keepNext")) is None:
                kn = ppr.makeelement(W("keepNext"), {})
                st = ppr.find(W("pStyle"))
                (st.addnext(kn) if st is not None else ppr.insert(0, kn))
    # Contribution 4 in Section I (AW, 2026-10-23), formatted like items
    # 1-3: a bold lead phrase, then the rest in the body font.
    item3 = find_para("Transparent baselines and error analysis:")
    tr3 = text_runs(item3)
    if len(tr3) != 2 or not ptext(item3).startswith(
            "Transparent baselines and error analysis:"):
        die("contribution item 3 not as expected")
    item4 = copy.deepcopy(item3)
    tr4 = text_runs(item4)
    set_text(tr4[0][1], "An extended evaluation")
    set_text(tr4[1][1], (
        f" on {v['ind_n']} tickets written and cross-labelled by two authors "
        f"(κ = {v['k_cat']:.3f}) and {v['rw_n']} real tickets from a public "
        "dataset, reported with its negative result."))
    item3.addnext(item4)
    # New appendices D, E, F after Appendix C
    for el in ([appendix_heading("Appendix D: Model Configurations and "
                                 "Hyperparameters")] + block_iv
               + [appendix_heading("Appendix E: Verification Checks")]
               + block_ii
               + [appendix_heading("Appendix F: Methodological Lessons"),
                  lessons]):
        anchor.addnext(el)
        anchor = el

    # 9. final checks ---------------------------------------------------------
    full = "".join(t.text or "" for t in body.iter(W("t")))
    left = sorted(set(marker_re.findall(full)))
    shaded = [r for r in body.iter(W("r")) if is_marker_run(r)]
    if left or shaded or "[[FILL" in full:
        die(f"unfilled markers {left}; {len(shaded)} shaded run(s) remain")
    for ref in ("[ref]", "{v[", "{'", "??", "Error! Reference"):
        if ref in full:
            die(f"leftover placeholder or broken reference {ref!r}")
    # Every cited table, figure, appendix and section must exist.
    texts = [ptext(p) for p in all_paras()]
    tables = {t.strip()[6:] for t in texts if re.fullmatch(r"TABLE [IVX]+", t.strip())}
    figs = {m.group(1) for t in texts for m in [re.match(r"Fig\. (\d+)\. ", t)] if m}
    apps = {m.group(1) for t in texts for m in [re.match(r"Appendix ([A-Z]): ", t)] if m}
    secs, cur = set(), None
    for t in texts:
        m = re.match(r"(I|II|III|IV|V|VI|VII|VIII|IX|X)\. [A-Z]", t)
        if m:
            cur = m.group(1)
            secs.add(cur)
        m = re.match(r"([A-H])\. [A-Z]", t)
        if m and cur and len(t) < 80:
            secs.add(f"{cur}-{m.group(1)}")
    cited = {
        "Table": (set(re.findall(r"Tables? ([IVX]+)\b", full)), tables),
        "Fig.": (set(re.findall(r"Fig\. (\d+)", full)), figs),
        "Appendix": (set(re.findall(r"Appendix ([A-Z])\b", full)), apps),
        "Section": (set(re.findall(r"Section ([IVX]+(?:-[A-H])?)\b", full)),
                    secs),
    }
    for kind, (refs, have) in cited.items():
        missing = sorted(refs - have)
        if missing:
            die(f"{kind} cited but not present: {missing} (have {sorted(have)})")
    print("[refs] " + "; ".join(f"{k} {sorted(r)}" for k, (r, _) in cited.items()))
    d.save(V4)
    print(f"[write] {os.path.relpath(V4, PROJECT_ROOT)}")


if __name__ == "__main__":
    main()
